"""Datos de evaluación reproducibles para Service Desk y CMDB."""
from __future__ import annotations

import random
from datetime import timedelta
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (ArticuloConocimiento, ClasificacionCI, Cola, ColaMiembro,
    DispositivoCI, ServicioCI, ServicioNivel2, Ticket, TicketCI, TicketHistorial, Usuario,
    UsuarioRol, utcnow)
from app.security import hash_password
from app.service_desk import PROCESOS, TIPOS


DEMO_USERS = [
    ("Solicitante Demo", "solicitante", "Solicitante-Demo-2026!", "solicitante"),
    ("Agente Demo", "agente", "Agente-Demo-2026!", "agente"),
    ("Aprobador Demo", "aprobador", "Aprobador-Demo-2026!", "aprobador"),
]


def seed_service_desk(db: Session) -> dict:
    users = {}
    for name, login, password, operational_role in DEMO_USERS:
        user = db.scalar(select(Usuario).where(Usuario.login == login))
        if not user:
            user = Usuario(nombre=name, login=login, password_hash=hash_password(password), rol="consulta")
            db.add(user); db.flush()
        users[operational_role] = user
        if not db.scalar(select(UsuarioRol).where(UsuarioRol.usuario_id == user.id, UsuarioRol.rol == operational_role)):
            db.add(UsuarioRol(usuario_id=user.id, rol=operational_role))
    queues = {}
    for i, process in enumerate(PROCESOS, 1):
        q = db.scalar(select(Cola).where(Cola.codigo == f"Q-{i:02d}"))
        if not q:
            q = Cola(codigo=f"Q-{i:02d}", nombre=process.replace("_", " ").title(),
                descripcion=f"Cola demostrativa para {process}", proceso=process,
                nivel_soporte=1 if process != "problemas" else 3,
                objetivo_primera_min=60, objetivo_resolucion_min=1440)
            db.add(q); db.flush()
        queues[process] = q
        for role, perms in [("agente", (True,True,True,False)), ("aprobador", (True,True,True,True))]:
            u=users[role]
            m=db.scalar(select(ColaMiembro).where(ColaMiembro.cola_id==q.id,ColaMiembro.usuario_id==u.id))
            if not m:
                db.add(ColaMiembro(cola_id=q.id,usuario_id=u.id,puede_consultar=perms[0],puede_atender=perms[1],puede_transferir=perms[2],puede_aprobar=perms[3]))
    if not db.scalar(select(ArticuloConocimiento.id)):
        db.add_all([
            ArticuloConocimiento(titulo="Restablecer una contraseña",contenido="Use el portal de acceso, valide su identidad y solicite el restablecimiento.",etiquetas="acceso,nivel-0"),
            ArticuloConocimiento(titulo="Antes de reportar un incidente",contenido="Registre servicio, síntomas, hora, impacto y cualquier CI conocido.",etiquetas="incidente,nivel-0"),
        ])
    db.commit()
    return {"usuarios": len(users), "colas": len(queues)}


def ensure_cmdb_hierarchy(db: Session, source: Path = Path("CMDB.xlsx")) -> list[ClasificacionCI]:
    # Se valida la fuente, pero nunca se escribe en ella.
    source_types = ["Hardware", "Software"]
    if source.is_file():
        ws = load_workbook(source, read_only=True, data_only=True)["CMDB"]
        found = {str(ws.cell(r,3).value or "").strip() for r in range(5,ws.max_row+1)} - {""}
        if found: source_types = sorted(found)
    root = db.scalar(select(ClasificacionCI).where(ClasificacionCI.codigo=="CI"))
    if not root:
        root=ClasificacionCI(codigo="CI",nombre="Elemento de configuración",origen="CMDB.xlsx!CMDB:C4:C6; nivel raíz de navegación")
        db.add(root);db.flush()
    result=[]
    definitions={
        "Hardware":{"numero_serie":"texto","ip":"texto opcional","cpu":"texto","memoria_gb":"entero"},
        "Software":{"version":"texto","licencia":"texto opcional","servidor":"texto opcional"},
    }
    for name in source_types:
        code=name.upper().replace(" ","_")
        x=db.scalar(select(ClasificacionCI).where(ClasificacionCI.padre_id==root.id,ClasificacionCI.codigo==code))
        if not x:
            x=ClasificacionCI(codigo=code,nombre=name,padre_id=root.id,atributos_definidos=definitions.get(name,{}),origen="CMDB.xlsx!CMDB:C5:C6")
            db.add(x);db.flush()
        result.append(x)
    db.commit();return result


def generate_cmdb(db: Session, count: int = 5000, seed: int = 20261028) -> dict:
    if count < 5000: raise ValueError("La carga de evaluación exige al menos 5000 dispositivos")
    classes=ensure_cmdb_hierarchy(db); rng=random.Random(seed)
    users=list(db.scalars(select(Usuario).where(Usuario.activo.is_(True)).order_by(Usuario.id)))
    locations=["Centro de datos A","Centro de datos B","Oficina Norte","Oficina Sur","Nube privada"]
    existing=set(db.scalars(select(DispositivoCI.etiqueta).where(DispositivoCI.etiqueta.like(f"SYN-{seed}-%"))))
    created=0
    for i in range(1,count+1):
        tag=f"SYN-{seed}-{i:05d}"
        if tag in existing: continue
        cls=classes[(i-1)%len(classes)]; hardware=cls.nombre.strip().lower()=="hardware"
        attrs=({"cpu":rng.choice(["4 cores","8 cores","16 cores"]),"memoria_gb":rng.choice([8,16,32,64]),"ip":f"10.{(i//65536)%255}.{(i//256)%255}.{i%255}"}
            if hardware else {"editor":rng.choice(["Open Source","Proveedor A","Proveedor B"]),"entorno":rng.choice(["producción","pruebas","desarrollo"])})
        db.add(DispositivoCI(etiqueta=tag,nombre=("Equipo" if hardware else "Aplicación")+f" {i:05d}",clasificacion_id=cls.id,
            estado=rng.choice(["Activo","Activo","Activo","Mantenimiento"]),ubicacion=locations[i%len(locations)],
            responsable_id=users[i%len(users)].id if users else None,numero_serie=f"SN-{seed}-{i:05d}" if hardware else None,
            version=None if hardware else f"{1+i%5}.{i%10}",nivel_seguridad=rng.choice(["Bajo","Medio","Alto"]),
            cumplimiento="Cumple",estado_configuracion="Aprobado",licencia=None if hardware else f"LIC-{seed}-{i:05d}",atributos=attrs))
        created+=1
        if created%500==0:db.flush()
    db.commit()
    totals={name:int(n) for name,n in db.execute(select(ClasificacionCI.nombre,func.count(DispositivoCI.id)).join(DispositivoCI).group_by(ClasificacionCI.nombre))}
    return {"semilla":seed,"solicitados":count,"creados":created,"total_con_semilla":db.scalar(select(func.count()).select_from(DispositivoCI).where(DispositivoCI.etiqueta.like(f"SYN-{seed}-%"))),"por_categoria":totals}


def seed_tickets(db: Session) -> dict:
    services=list(db.scalars(select(ServicioNivel2).where(ServicioNivel2.activo.is_(True)).order_by(ServicioNivel2.id)))
    if not services:return {"creados":0,"advertencia":"Importe el catálogo primero"}
    requester=db.scalar(select(Usuario).where(Usuario.login=="solicitante"));agent=db.scalar(select(Usuario).where(Usuario.login=="agente"))
    cis=list(db.scalars(select(DispositivoCI).order_by(DispositivoCI.id).limit(35)))
    created=0
    incident_categories=[f"I{i}" for i in range(7)];event_categories=["E0","E1","E2","E3"]
    for process,flow in PROCESOS.items():
        queue=db.scalar(select(Cola).where(Cola.proceso==process,Cola.activo.is_(True)).order_by(Cola.id))
        target = 7 if process == "incidentes" else 5
        existing = db.scalar(select(func.count()).select_from(Ticket).where(Ticket.proceso==process, Ticket.titulo.like(f"Caso demo {process} %"))) or 0
        for j in range(existing, target):
            data={}
            if process=="incidentes":data={"categoria":incident_categories[j%7],"interrupcion":j>0,"redundancia":"disponible" if j%2 else "no_disponible","justificacion_categoria":"Caso de demostración"}
            if process=="eventos":data={"categoria":event_categories[j%4],"evaluacion":"Caso demostrativo"}
            if process=="cambios":data={"impacto":"medio","riesgo":"bajo","plan_pruebas":"Pruebas funcionales","plan_reversion":"Restaurar versión anterior"}
            if process=="accesos":data={"beneficiario_id":requester.id,"recurso":"Portal","permiso":"lectura"}
            state=flow[min(j,len(flow)-2)]
            t=Ticket(numero="PENDIENTE",titulo=f"Caso demo {process} {j+1}",descripcion="Ticket sintético, sin datos personales reales.",solicitante_id=requester.id,creador_id=requester.id,
                servicio_id=services[created%len(services)].id,tipo=TIPOS[process],proceso=process,cola_id=queue.id,responsable_id=agent.id if j else None,
                estado=state,prioridad=["baja","media","alta","critica","media"][j%5],impacto="medio",urgencia="media",canal="portal",
                sin_ci_justificacion=None if cis else "No aplica CI en el entorno vacío",datos_proceso=data,
                primera_atencion_en=utcnow()-timedelta(minutes=30) if j else None)
            db.add(t);db.flush();t.numero=f"SD-{t.id:07d}"
            if cis: db.add(TicketCI(ticket_id=t.id,ci_id=cis[created%len(cis)].id))
            db.add(TicketHistorial(ticket_id=t.id,actor_id=requester.id,accion="creacion_demo",nuevo={"estado":state}));created+=1
    # Normaliza los siete casos de incidente por posición. Esto también corrige
    # conjuntos demo creados por versiones anteriores sin borrar historial.
    demo_incidents=list(db.scalars(select(Ticket).where(Ticket.proceso=="incidentes",Ticket.titulo.like("Caso demo incidentes %")).order_by(Ticket.titulo)))
    for i,ticket in enumerate(demo_incidents[:7]):
        payload=dict(ticket.datos_proceso or {});payload["categoria"]=incident_categories[i]
        payload["seguridad"] = i == 6
        payload.setdefault("justificacion_categoria","Caso de demostración")
        ticket.datos_proceso=payload
    db.commit();return {"creados":created,"total":db.scalar(select(func.count()).select_from(Ticket))}
