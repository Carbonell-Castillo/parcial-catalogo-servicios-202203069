from __future__ import annotations

import hashlib
from dataclasses import dataclass, asdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from app.models import (
    ClaseServicio, Criticidad, ImportacionLote, ImportacionObservacion,
    ServicioNivel1, ServicioNivel2, TipoServicio, utcnow,
)

SHEET = "Servicios Externos"
CANONICAL_SE12 = "Suministrar Analitica"

@dataclass
class ImportSummary:
    lote_id: int
    creados: int = 0
    actualizados: int = 0
    omitidos: int = 0
    observados: int = 0
    nivel1: int = 0
    nivel2: int = 0

    def dict(self):
        return asdict(self)

def text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None

def decimal_or_none(value: Any) -> Decimal | None:
    if value is None or text(value) is None:
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None

def merged_map(ws) -> dict[str, tuple[str, str]]:
    result = {}
    for merged in ws.merged_cells.ranges:
        anchor = ws.cell(merged.min_row, merged.min_col).coordinate
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                result[ws.cell(row, col).coordinate] = (anchor, str(merged))
    return result

def resolved(ws, row: int, col: int, merges: dict[str, tuple[str, str]]):
    cell = ws.cell(row, col)
    origin = merges.get(cell.coordinate)
    if origin:
        return ws[origin[0]].value, origin[1]
    return cell.value, cell.coordinate

def observe(db: Session, lote_id: int, row: int | None, code: str | None, kind: str, detail: str,
            original: Any = None, canonical: Any = None, source_range: str | None = None):
    db.add(ImportacionObservacion(lote_id=lote_id, hoja=SHEET, fila=row, rango_origen=source_range,
        codigo=code, tipo=kind, detalle=detail, valor_original=text(original), valor_canonico=text(canonical)))

def catalog_id(db: Session, model, value: Any, lote_id: int, row: int) -> int | None:
    name = text(value)
    if not name:
        return None
    item = db.scalar(select(model).where(func.lower(model.nombre) == name.lower()))
    if item:
        return item.id
    observe(db, lote_id, row, None, "catalogo_desconocido", f"Valor no incluido en catálogo controlado: {name}", name)
    return None

def import_excel(db: Session, path: str | Path) -> ImportSummary:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"No existe el archivo original: {source}")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    lote = ImportacionLote(archivo=source.name, sha256=digest)
    db.add(lote)
    db.flush()
    summary = ImportSummary(lote_id=lote.id)
    try:
        book = load_workbook(source, data_only=True, read_only=False)
        if SHEET not in book.sheetnames:
            raise ValueError(f"No existe la hoja requerida: {SHEET}")
        ws = book[SHEET]
        merges = merged_map(ws)
        n1_names: dict[str, str] = {}
        seen_n1: set[str] = set()
        seen_n2: set[str] = set()
        for row in range(5, 102):
            vals, ranges = [], []
            for col in range(1, 13):
                value, source_range = resolved(ws, row, col, merges)
                vals.append(value)
                ranges.append(source_range)
            code_cell = ws.cell(row, 3).coordinate
            code_merge = merges.get(code_cell)
            if code_merge and code_merge[0] != code_cell:
                # La fila pertenece al mismo servicio N2 cuyo código está en el
                # ancla combinada. Se distingue como continuación y no genera
                # un segundo registro ni sobrescribe atributos con vacíos.
                summary.omitidos += 1
                continue
            n1_code, n1_name, n2_code, n2_name = map(text, vals[:4])
            if not n2_code:
                if any(text(v) for v in vals[3:]):
                    observe(db, lote.id, row, n1_code, "fila_sin_codigo_n2",
                            "Fila con datos pero sin código N2; no fue asociada automáticamente", source_range=f"A{row}:L{row}")
                    summary.observados += 1
                    summary.omitidos += 1
                else:
                    summary.omitidos += 1
                continue
            if row == 101 and n2_code == "SE.12.3" and not n1_code and not n1_name:
                # Excepción documentada del archivo real: esta fila no forma
                # parte de una combinación A/B, pero su código explícito la
                # identifica inequívocamente como hija de SE.12.
                n1_code, n1_name = "SE.12", CANONICAL_SE12
                observe(db, lote.id, row, n2_code, "padre_inferido_se12",
                        "Fila 101 sin A/B: se vinculó SE.12.3 al padre SE.12 por regla explícita",
                        None, "SE.12", f"A{row}:B{row}")
                summary.observados += 1
            if not n1_code or not n1_name or not n2_name:
                observe(db, lote.id, row, n2_code, "campo_obligatorio_ausente",
                        "Faltan código/nombre N1 o nombre N2; fila omitida", source_range=f"A{row}:L{row}")
                summary.observados += 1
                summary.omitidos += 1
                continue
            canonical_n1 = n1_name
            if n1_code == "SE.12":
                canonical_n1 = CANONICAL_SE12
                if n1_name != CANONICAL_SE12:
                    observe(db, lote.id, row, n1_code, "conflicto_nombre_n1",
                            "SE.12 usa el primer nombre de la fila 99 como canónico; se conserva el alternativo",
                            n1_name, CANONICAL_SE12, ranges[1])
                    summary.observados += 1
            previous = n1_names.setdefault(n1_code, canonical_n1)
            if previous != canonical_n1:
                observe(db, lote.id, row, n1_code, "conflicto_nombre_n1",
                        "Un código N1 presenta nombres distintos; se conserva el canónico", n1_name, previous, ranges[1])
                summary.observados += 1
                canonical_n1 = previous
            level1 = db.scalar(select(ServicioNivel1).where(ServicioNivel1.codigo == n1_code))
            if not level1:
                level1 = ServicioNivel1(codigo=n1_code, nombre=canonical_n1, origen_hoja=SHEET, origen_fila=row, origen_rango=f"{ranges[0]},{ranges[1]}")
                db.add(level1); db.flush(); summary.creados += 1
            elif level1.nombre != canonical_n1:
                level1.nombre = canonical_n1; summary.actualizados += 1
            seen_n1.add(n1_code)
            incomplete = row >= 99 or any(vals[i] is None for i in (4, 5, 6, 7, 9))
            data = dict(
                codigo_original=n2_code, nombre=n2_name, nivel1_id=level1.id,
                indicador_activo_excel=text(vals[4]),
                clase_id=catalog_id(db, ClaseServicio, vals[5], lote.id, row),
                criticidad_id=catalog_id(db, Criticidad, vals[6], lote.id, row),
                tipo_id=catalog_id(db, TipoServicio, vals[7], lote.id, row),
                descripcion=text(vals[8]), metrica=text(vals[9]),
                minimo=decimal_or_none(vals[10]), maximo=decimal_or_none(vals[11]),
                estado_revision="pendiente" if incomplete else "completo",
                origen_hoja=SHEET, origen_fila=row,
                origen_rango=",".join(dict.fromkeys(ranges[:4])),
            )
            for index, field in ((10, "mínimo"), (11, "máximo")):
                if text(vals[index]) is not None and decimal_or_none(vals[index]) is None:
                    observe(db, lote.id, row, n2_code, "numero_invalido",
                            f"El valor de {field} no es numérico; se conservó como desconocido", vals[index], None, ranges[index])
                    data["estado_revision"] = "pendiente"; summary.observados += 1
            if data["minimo"] is not None and data["maximo"] is not None and data["minimo"] > data["maximo"]:
                observe(db, lote.id, row, n2_code, "rango_invalido", "Mínimo mayor que máximo; ambos se conservaron como NULL",
                        f"{data['minimo']}..{data['maximo']}", None, f"K{row}:L{row}")
                data["minimo"] = data["maximo"] = None; data["estado_revision"] = "pendiente"; summary.observados += 1
            service = db.scalar(select(ServicioNivel2).where(ServicioNivel2.codigo == n2_code))
            if not service:
                db.add(ServicioNivel2(codigo=n2_code, **data)); summary.creados += 1
            elif n2_code in seen_n2:
                if service.nombre != n2_name:
                    observe(db, lote.id, row, n2_code, "conflicto_atributos_n2",
                            "Fila repetida con nombre diferente; se conserva el primero", n2_name, service.nombre, ranges[3])
                    summary.observados += 1
                summary.omitidos += 1
            else:
                changed = any(getattr(service, key) != value for key, value in data.items())
                if changed:
                    for key, value in data.items(): setattr(service, key, value)
                    summary.actualizados += 1
                else:
                    summary.omitidos += 1
            seen_n2.add(n2_code)
        db.flush()
        summary.nivel1 = len(seen_n1)
        summary.nivel2 = len(seen_n2)
        if summary.nivel1 != 12 or summary.nivel2 != 46:
            observe(db, lote.id, None, None, "control_conteo",
                    f"Conteo obtenido {summary.nivel1}/{summary.nivel2} N1/N2; esperado 12/46")
            summary.observados += 1
        lote.creados, lote.actualizados, lote.omitidos, lote.observados = summary.creados, summary.actualizados, summary.omitidos, summary.observados
        lote.finalizada_en, lote.estado = utcnow(), "completada"
        db.commit()
        return summary
    except Exception:
        db.rollback()
        raise
