from pathlib import Path
import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from app.importer import import_excel
from app.models import ImportacionObservacion, ServicioNivel1, ServicioNivel2

def make_synthetic_catalog(path: Path):
    """Fixture controlado: prueba reglas, no pretende reemplazar el Excel original."""
    wb=Workbook(); ws=wb.active; ws.title="Servicios Externos"
    headers=["COD.N1","SERVICIO - Nivel 1","COD.N2","SERVICIO - Nivel 2","ACTIVO","CLASE DE SERVICIO","CRITICIDAD","TIPO DE SERVICIO","Descripción","Métrica","Minimo","Maximo"]
    for c,v in enumerate(headers,1): ws.cell(4,c,v)
    rows=list(range(5,48))+[99,100,101]
    blocks=[]
    for idx,row in enumerate(rows):
        if idx<43:
            group=min(idx//4+1,11)
            n1=f"SE.{group}"; n1name=f"Servicio nivel uno {group}"
            values=[n1,n1name,f"SE.{group}.{idx+1}",f"Servicio {idx+1}","S","A DEMANDA","Normal","Back End","Descripción","Porcentaje",0,100]
        else:
            sub=idx-42; n1="SE.12"; n1name="Suministrar Analitica" if row==99 else "Mantener Tableros de Control"
            values=[None if row==101 else n1,None if row==101 else n1name,f"SE.12.{sub}",f"Servicio incompleto {sub}",None,None,None,None,None,None,None,None]
        for col,value in enumerate(values,1): ws.cell(row,col,value)
    # Las primeras cuatro filas comparten N1 mediante combinaciones; verifica la recuperación desde el ancla.
    ws.merge_cells("A5:A8"); ws.merge_cells("B5:B8")
    wb.save(path)

def test_p06_p07_p08_synthetic_import_is_idempotent_and_traceable(db,tmp_path):
    source=tmp_path/"catalogo-sintetico.xlsx"; make_synthetic_catalog(source)
    first=import_excel(db,source)
    assert first.nivel1==12 and first.nivel2==46
    assert db.scalar(select(func.count()).select_from(ServicioNivel2))==46
    second=import_excel(db,source)
    assert second.creados==0
    assert db.scalar(select(func.count()).select_from(ServicioNivel2))==46
    se12=db.scalar(select(ServicioNivel1).where(ServicioNivel1.codigo=="SE.12"))
    assert se12.nombre=="Suministrar Analitica"
    missing=db.scalar(select(ServicioNivel2).where(ServicioNivel2.codigo=="SE.12.1"))
    assert missing.estado_revision=="pendiente" and missing.clase_id is None and missing.minimo is None
    observations=db.scalars(select(ImportacionObservacion).where(ImportacionObservacion.codigo=="SE.12")).all()
    assert any(x.valor_original=="Mantener Tableros de Control" and x.valor_canonico=="Suministrar Analitica" for x in observations)

def test_merged_n2_continuation_creates_only_anchor_service(db,tmp_path):
    source=tmp_path/"continuacion.xlsx"
    wb=Workbook(); ws=wb.active; ws.title="Servicios Externos"
    values=["SE.01","Infraestructura","SE.01.01","Servicio combinado","S","A DEMANDA","Normal","Back End",None,"Métrica",1,2]
    for col,value in enumerate(values,1): ws.cell(5,col,value)
    ws.merge_cells("A5:A7"); ws.merge_cells("B5:B7"); ws.merge_cells("C5:C7"); ws.merge_cells("D5:D7")
    for row in (6,7):
        for col,value in enumerate(["S","A DEMANDA","Normal","Back End"],5): ws.cell(row,col,value)
    wb.save(source)
    result=import_excel(db,source)
    assert result.nivel2==1
    assert db.scalar(select(func.count()).select_from(ServicioNivel2))==1
    service=db.scalar(select(ServicioNivel2))
    assert service.origen_rango and "C5:C7" in service.origen_rango

@pytest.mark.skipif(not Path("data/CatalogoServicios.xlsx").is_file(),reason="El archivo original no fue proporcionado")
def test_original_excel_has_required_counts(db):
    summary=import_excel(db,"data/CatalogoServicios.xlsx")
    assert (summary.nivel1,summary.nivel2)==(12,46)
